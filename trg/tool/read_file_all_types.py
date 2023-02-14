import os
import subprocess
import logging
from trg.tool.jsontool import read_json_



class ReadFile:
    def read_file(self, filename):
        # this is only used linux system
        '''
        buff = subprocess.Popen("cat '%s'" %filename, stdout=subprocess.PIPE, shell=True)
        content = buff.communicate()[0]'''

        try:
            with open(filename, 'r', encoding='utf-8') as f:
                content = f.read()
        except:
            with open(filename, 'r', encoding='Latin_1') as f:
                content = f.read()
        return content


def read_file(filename):
    # this is only used linux system
    '''
    buff = subprocess.Popen("cat '%s'" %filename, stdout=subprocess.PIPE, shell=True)
    content = buff.communicate()[0]'''

    try:
        with open(filename,'r', encoding='utf-8') as f:
            content = f.read()
    except:
        with open(filename,'r', encoding='Latin_1') as f:
            content = f.read()
    return content


def read_file_xls_(filename):
    import xlrd
    try:
        content = xlrd.open_workbook(filename)
        tables = content.sheets()
        sheets_cnt = 1
        for table in tables:
            for row in range(0, table.nrows):
                tmp = ''
                for col in range(0, table.ncols):
                    if table.cell(row, col).value != None:
                        tmp += str(table.cell(row,col).value) + '\t'
            sheets_cnt += 1
    except Exception(e):
        print ("read xls file error,watch log file")
        logging.debug("read xls file error,because %s",str(e))

'''
def read_file_doc_(filename):
    from docx import Document
#    from document import Document
#    from docx import Document
    content =[]
    try:
        document = Document(filename)
        for page in document.paragraphs:
            content.append(page.text)
    except Exception,e:
        print "read doc file error,watch log file"
        logging.debug("read doc file error,because %s",str(e))

    return content
'''

def read_file_pdf_(filename):

    from pdfrw import PdfReader
    content = []
    try:
        cnt = PdfReader(filename)
        for page in cnt.pages:
            content.append(page.Contents)
    except Exception(e):
        print ("read pdf file error,watch log file")
        logging.debug("read pdf file error,because %s",str(e))

    return content


def read_file_txt_(filename):

    content = ""
    try:
        with open(filename,'r', encoding='utf-8') as f:
            for line in f.readlines():
                if not len(line) or line.startswith('#'):
                    continue
                content += line
    except:
        with open(filename,'r', encoding='Latin_1') as f:
            for line in f.readlines():
                if not len(line) or line.startswith('#'):
                    continue
                content += line
"""
    try:
        for line in f.readlines():
            if not len(line) or line.startswith('#'):
                continue
            content += line
    except Exception(e):
        print ("read txt file error, watch log file")
        logging.debug("read txt file error,because %s", str(e))

    return content
"""


def read_file_all_types(path_):
    # depend on the type of file, to decide use which particular function
    # the control of read file.
    content = None
    badname = ['swp', 'jpg', 'gif', 'png', 'pyc', 'py.swp', 'class']
    deps = os.path.splitext(path_)[1][1:]

    if deps =='xls' or deps == 'xlsx':
        content = read_file_xls_(path_)
#    elif deps == 'doc' or deps == 'docx':
#        content = read_file_doc_(path_)
    elif deps == 'pdf':
        content = read_file_pdf_(path_)
    elif deps == 'txt':
        content = read_file_txt_(path_)
    elif deps == 'json':
        content = read_json_(path_)
    elif deps not in badname:
        content = read_file(path_)
    return content
