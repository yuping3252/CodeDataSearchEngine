#!/usr/bin/python
import os
import socket
import ssl
import PyPDF2
import docx
import docx2txt

from PyQt4.QtCore import *
from PyQt4.QtGui  import *


class RightLoadFile:
    def __init__(self, rightpane):
        self.rightpane = rightpane


    def load_file(self):

        fdialog = QFileDialog()

        fname_full = fdialog.getOpenFileName()
        print("file name full ", fname_full)

        fname_lst  = fname_full.split('.')

        fsuffix = fname_lst[len(fname_lst)-1]
        fname   = fname_full[:len(fname_full) - len(fsuffix) - 1]

        fobj = ""

        if fsuffix == 'pdf':
            print("pdf file")
            fobj = open(fname_full, 'rb')
            print("pdf file, ............................ loaded")
            #pdfReader = PyPDF2.PdfFileReader(fobj)
            #pageObj = pdfReader.getPage(0)
            #pageText = pageObj.extractText()
            #print("pageText=", pageText)

        elif fsuffix == 'xlsx':
            print("xlsx file")
            fobj = openpyxl.load_workbook(fname_full)
            print("xlsx file ............................... loaded")

        elif fsuffix == 'doc':
            print("doc file")
            fobj = docx.Document(fname_full)
            print("doc file, ............................... loaded ")

        elif fsuffix == 'docx':
            print("docx file")
            my_text = docx2txt.process(fname_full)
            print("docx file, ............................... loaded, text=", my_text)

            fobj = docx.Document(fname_full)
            print("docx file, ............................... loaded ")

        elif fsuffix == 'txt' or fsuffix == 'xml' or fsuffix == 'py' or fsuffix == 'sql':
            print("text file")
            fobj = open(fname_full, 'r')
            print("text file, ............................... loaded")

        elif fsuffix == 'exe':
            print("exe file")
            fobj = open(fname_full, 'rb')
            print("exe file, ............................... loaded")

        file_comb = [fname_full, fname, fsuffix, fobj]

        self.rightpane.add_load_file_icon(file_comb)




